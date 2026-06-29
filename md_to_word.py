#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def parse_markdown(md_text):
    lines = md_text.split('\n')
    elements = []
    in_code_block = False
    code_content = []
    code_lang = ''
    in_table = False
    table_rows = []

    for line in lines:
        # 代码块处理
        if line.strip().startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = line.strip().replace('```', '').strip()
                code_content = []
            else:
                elements.append(('code', '\n'.join(code_content)))
                in_code_block = False
            continue

        if in_code_block:
            code_content.append(line)
            continue

        # 表格处理
        if '|' in line and line.strip().startswith('|') and line.strip().endswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            # 跳过分隔线
            if re.match(r'^\|[\s\-:]+\|$', line.strip()):
                continue
            cells = [c.strip() for c in line.strip().split('|')[1:-1]]
            table_rows.append(cells)
            continue
        else:
            if in_table and table_rows:
                elements.append(('table', table_rows))
                in_table = False
                table_rows = []

        if not line.strip():
            elements.append(('blank', ''))
            continue

        # 标题
        h_match = re.match(r'^(#{1,6})\s+(.+)$', line)
        if h_match:
            level = len(h_match.group(1))
            text = h_match.group(2)
            elements.append(('heading', level, text))
            continue

        # 分隔线
        if re.match(r'^---+$', line.strip()):
            elements.append(('hr', ''))
            continue

        # 列表项
        li_match = re.match(r'^(\s*)-\s+(.+)$', line)
        if li_match:
            text = li_match.group(2)
            elements.append(('list_item', text))
            continue

        ol_match = re.match(r'^(\s*)\d+\.\s+(.+)$', line)
        if ol_match:
            text = ol_match.group(2)
            elements.append(('ordered_item', text))
            continue

        # 普通段落
        elements.append(('paragraph', line))

    if in_table and table_rows:
        elements.append(('table', table_rows))

    return elements


def add_inline_text(paragraph, text):
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            paragraph.add_run(part)


def md_to_word(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    elements = parse_markdown(md_text)

    doc = Document()

    # 设置默认中文字体
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    doc.styles['Normal'].font.size = Pt(12)

    # 页面设置
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    for elem in elements:
        etype = elem[0]

        if etype == 'blank':
            doc.add_paragraph('')
            continue

        if etype == 'heading':
            level = elem[1]
            text = elem[2]
            if level == 1:
                p = doc.add_heading('', level=0)
                run = p.add_run(text)
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run.font.size = Pt(22)
                run.font.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif level == 2:
                p = doc.add_heading('', level=1)
                run = p.add_run(text)
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run.font.size = Pt(16)
                run.font.bold = True
            elif level == 3:
                p = doc.add_heading('', level=2)
                run = p.add_run(text)
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run.font.size = Pt(14)
                run.font.bold = True
            else:
                p = doc.add_heading('', level=min(level, 3))
                run = p.add_run(text)
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run.font.size = Pt(12)
                run.font.bold = True
            continue

        if etype == 'paragraph':
            text = elem[1]
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(24)
            p.paragraph_format.line_spacing = 1.5
            add_inline_text(p, text)
            continue

        if etype == 'list_item':
            text = elem[1]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing = 1.5
            add_inline_text(p, text)
            continue

        if etype == 'ordered_item':
            text = elem[1]
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.line_spacing = 1.5
            add_inline_text(p, text)
            continue

        if etype == 'code':
            code_text = elem[1]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            p.paragraph_format.line_spacing = 1.2
            continue

        if etype == 'table':
            rows = elem[1]
            if not rows:
                continue
            num_cols = len(rows[0])
            table = doc.add_table(rows=len(rows), cols=num_cols)
            table.style = 'Light Grid Accent 1'
            for i, row in enumerate(rows):
                for j, cell_text in enumerate(row):
                    if j >= num_cols:
                        continue
                    cell = table.rows[i].cells[j]
                    cell.text = ''
                    p = cell.paragraphs[0]
                    if i == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p.add_run(cell_text)
                        run.bold = True
                    else:
                        add_inline_text(p, cell_text)
            doc.add_paragraph('')
            continue

        if etype == 'hr':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run('—' * 30)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            continue

    doc.save(docx_path)
    print(f'Word文档已生成: {docx_path}')


if __name__ == '__main__':
    md_file = '/workspace/房屋租赁信息管理系统-课程设计报告.md'
    docx_file = '/workspace/房屋租赁信息管理系统-课程设计报告.docx'
    md_to_word(md_file, docx_file)
